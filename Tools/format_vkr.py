from __future__ import annotations

import copy
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(os.environ["VKR_SOURCE"])
OUTPUT = Path(os.environ["VKR_OUTPUT"])
PAGE_COUNT = os.environ.get("VKR_PAGE_COUNT", "XX")

TNR = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, size: float = 14, bold: bool | None = None) -> None:
    run.font.name = TNR
    run._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def set_paragraph_font(paragraph: Paragraph, size: float = 14, bold: bool | None = None) -> None:
    if not paragraph.runs:
        run = paragraph.add_run()
        set_run_font(run, size, bold)
    for run in paragraph.runs:
        set_run_font(run, size, bold)


def set_normal_format(paragraph: Paragraph, first_line: bool = True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.first_line_indent = Cm(1.25) if first_line else Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_paragraph_font(paragraph, 14)


def set_compact_format(
    paragraph: Paragraph,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    size: float = 14,
    bold: bool | None = None,
    first_indent: float = 0,
    line_spacing: float = 1.0,
) -> None:
    paragraph.alignment = alignment
    pf = paragraph.paragraph_format
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.first_line_indent = Cm(first_indent)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = line_spacing
    set_paragraph_font(paragraph, size, bold)


def delete_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def paragraph_by_text(document: Document, text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def paragraph_startswith(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def insert_paragraph_before_element(document: Document, element, text: str = "") -> Paragraph:
    paragraph = document.add_paragraph(text)
    element.addprevious(paragraph._p)
    return paragraph


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = Paragraph(new_p, paragraph._parent)
    if text:
        result.add_run(text)
    return result


def add_page_field(paragraph: Paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])
    set_run_font(run, 14)


def add_toc_field(paragraph: Paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Оглавление обновляется в Microsoft Word"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    set_run_font(run, 14)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = TNR
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
    normal.font.size = Pt(14)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = TNR
        style._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
        style.font.size = Pt(14)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.keep_with_next = True


def configure_page(document: Document) -> None:
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)
        section.different_first_page_header_footer = True

        footer = section.footer
        for paragraph in footer.paragraphs:
            clear_paragraph(paragraph)
        footer_p = footer.paragraphs[0]
        set_compact_format(footer_p, WD_ALIGN_PARAGRAPH.CENTER, 14)
        add_page_field(footer_p)

        for first_page_part in (section.first_page_header, section.first_page_footer):
            for paragraph in first_page_part.paragraphs:
                clear_paragraph(paragraph)


def remove_front_matter(document: Document) -> Paragraph:
    intro = paragraph_by_text(document, "ВВЕДЕНИЕ")
    body = document._element.body
    for child in list(body):
        if child is intro._p:
            break
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return intro


def add_front_paragraph(document: Document, anchor: Paragraph, text: str = "") -> Paragraph:
    paragraph = document.add_paragraph(text)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def build_front_matter(document: Document, intro: Paragraph) -> None:
    created: list[Paragraph] = []

    def add(text: str = "", *, align=WD_ALIGN_PARAGRAPH.LEFT, bold=None, size=14, first=0, line=1.0):
        p = document.add_paragraph(text)
        set_compact_format(p, align, size, bold, first, line)
        created.append(p)
        return p

    for text in [
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ",
        "Федеральное государственное бюджетное образовательное учреждение",
        "высшего образования",
        "«КУБАНСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ»",
        "(ФГБОУ ВО «КубГУ»)",
    ]:
        add(text, align=WD_ALIGN_PARAGRAPH.CENTER)
    add()
    add("Факультет компьютерных технологий и прикладной математики", align=WD_ALIGN_PARAGRAPH.CENTER)
    add("Кафедра анализа данных и искусственного интеллекта", align=WD_ALIGN_PARAGRAPH.CENTER)
    add()
    add("Допустить к защите", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add("заведующий кафедрой", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add("д-р тех. наук, доцент", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add("____________ А. В. Коваленко", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add("«___» ____________ 2026 г.", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add()
    add("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add("(БАКАЛАВРСКАЯ РАБОТА)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add()
    add(
        "РАЗРАБОТКА МОБИЛЬНОГО ПРИЛОЖЕНИЯ-ЕЖЕДНЕВНИКА "
        "С ИСПОЛЬЗОВАНИЕМ МЕТОДА ГЕЙМИФИКАЦИИ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        line=1.5,
    )
    add()
    add("Работу выполнил(а)  __________________________  И. Д. Ерюшев")
    add("Направление подготовки 02.03.03 Математическое обеспечение и")
    add("администрирование информационных систем")
    add("Направленность (профиль) Технология программирования")
    add()
    add("Научный руководитель")
    add("канд. пед. наук, доцент  ______________________  В. А. Акиньшина")
    add()
    add("Нормоконтролер")
    add("канд. физ.-мат. наук, доцент  _________________  Г. В. Калайдина")
    add()
    add()
    add("Краснодар 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
    title_break = add()
    title_break.add_run().add_break(WD_BREAK.PAGE)

    ref_heading = add("РЕФЕРАТ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, line=1.5)
    ref_heading.style = document.styles["Heading 1"]
    add()
    for text in [
        f"Выпускная квалификационная работа содержит {PAGE_COUNT} с., 5 рис., 8 табл., 20 источников.",
        (
            "ГЕЙМИФИКАЦИЯ, UNITY, C#, NEXT.JS, TYPESCRIPT, FIREBASE, FIRESTORE, "
            "МОБИЛЬНОЕ ПРИЛОЖЕНИЕ, ВЕБ-ПРИЛОЖЕНИЕ, РОДИТЕЛЬСКИЙ КОНТРОЛЬ, "
            "ЭКСПЕРТНАЯ СИСТЕМА, АНАЛИТИКА, РЕКОМЕНДАЦИИ."
        ),
        (
            "Объектом исследования является процесс разработки клиент-серверной программной "
            "системы, предназначенной для мотивации ребенка к выполнению реальных повседневных "
            "задач средствами геймификации."
        ),
        (
            "Предмет исследования составляют архитектурные, программные и интерфейсные решения, "
            "обеспечивающие взаимодействие игрового Unity-клиента ребенка, родительского "
            "веб-приложения и облачного хранилища данных."
        ),
        (
            "Цель работы заключается в развитии прототипа мобильного приложения-ежедневника до "
            "целостной системы с синхронизацией заданий, адаптивным игровым интерфейсом, "
            "аналитикой выполнения и объяснимыми рекомендациями для родителя."
        ),
        (
            "В работе применены методы объектно-ориентированного проектирования, сравнительного "
            "анализа, прототипирования пользовательских интерфейсов, событийного моделирования "
            "и логического экспертного вывода."
        ),
        (
            "Результатом является программный комплекс, включающий Unity-клиент ребенка, "
            "родительскую веб-панель, общую модель данных Firestore, жизненный цикл задания, "
            "систему игровых наград, аналитические показатели и экспертный модуль рекомендаций. "
            "Геолокационные квесты обоснованы как направление дальнейшего развития."
        ),
    ]:
        p = add(text, line=1.5, first=1.25)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ref_break = add()
    ref_break.add_run().add_break(WD_BREAK.PAGE)

    toc_heading = add("СОДЕРЖАНИЕ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, line=1.5)
    toc_heading.style = document.styles["Heading 1"]
    toc_p = add(line=1.5)
    add_toc_field(toc_p)

    for p in created:
        intro._p.addprevious(p._p)


def normalize_headings(document: Document) -> None:
    structural = {
        "ВВЕДЕНИЕ",
        "ЗАКЛЮЧЕНИЕ",
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
        "РЕФЕРАТ",
        "СОДЕРЖАНИЕ",
    }
    chapter_titles = {
        "1. АНАЛИЗ ИСХОДНОГО ПРОТОТИПА И АРХИТЕКТУРНОЕ РАЗВИТИЕ СИСТЕМЫ":
            "1 Анализ исходного прототипа и архитектурное развитие системы",
        "2. РЕАЛИЗАЦИЯ УЛУЧШЕНИЙ И РЕЗУЛЬТАТЫ":
            "2 Реализация улучшений и результаты",
        "3. ПЕРСПЕКТИВЫ РАЗВИТИЯ СИСТЕМЫ: ГЕОЛОКАЦИОННЫЕ КВЕСТЫ, АНАЛИТИКА И ЭКСПЕРТНЫЕ РЕКОМЕНДАЦИИ":
            "3 Расширение системы: геолокационные квесты, аналитика и экспертные рекомендации",
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in chapter_titles:
            replace_paragraph_text(paragraph, chapter_titles[text])
            text = paragraph.text.strip()
        elif paragraph.style.name.startswith("Heading"):
            cleaned = re.sub(r"^(\d+(?:\.\d+)*)\.\s+", r"\1 ", text)
            if cleaned != text:
                replace_paragraph_text(paragraph, cleaned)
                text = cleaned

        if paragraph.style.name == "Heading 1":
            pf = paragraph.paragraph_format
            pf.page_break_before = text not in {"РЕФЕРАТ", "СОДЕРЖАНИЕ"}
            pf.keep_with_next = True
            if text in structural:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf.first_line_indent = Cm(0)
                set_paragraph_font(paragraph, 14, True)
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.first_line_indent = Cm(1.25)
                set_paragraph_font(paragraph, 14, True)
        elif paragraph.style.name in {"Heading 2", "Heading 3"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(1.25)
            paragraph.paragraph_format.keep_with_next = True
            set_paragraph_font(paragraph, 14, True)


def replace_exact(document: Document, old: str, new: str) -> None:
    paragraph = paragraph_by_text(document, old)
    replace_paragraph_text(paragraph, new)


def edit_content(document: Document) -> None:
    replacements = {
        (
            "Поэтому следующая стадия разработки затронула не только с улучшением внешнего вида, "
            "но и с переходом от отдельного игрового прототипа к распределенной системе. "
            "В проект была добавлена родительская веб-версия на Next.js и TypeScript, использующая "
            "Firebase Authentication и Firestore. Unity-клиент получил REST-мост к Firebase, "
            "модели профиля ребенка, состояния игры и реальных заданий. В итоге, ребенок "
            "взаимодействует с игрой, а родитель - с веб-панелью, при этом обе части работают "
            "с общей моделью данных."
        ): (
            "Поэтому следующая стадия разработки была связана не только с улучшением внешнего вида, "
            "но и с переходом от отдельного игрового прототипа к распределенной системе. "
            "В проект добавлено родительское веб-приложение на Next.js и TypeScript, использующее "
            "Firebase Authentication и Firestore. Unity-клиент получил REST-интерфейс доступа к "
            "Firebase, модели профиля ребенка, игрового состояния и реальных заданий. В результате "
            "ребенок взаимодействует с игрой, а родитель – с веб-панелью; обе части используют "
            "общую модель данных."
        ),
        (
            "–\tНастроить модель синхронизации заданий и игрового состояния через Firebase "
            "Authentication, Firestore и Cloud Functions."
        ): (
            "– настроить синхронизацию заданий и игрового состояния через Firebase Authentication "
            "и Firestore, а также подготовить Cloud Functions для доверенных серверных операций;"
        ),
        (
            "-\tРазработать подход к созданию экспертной рекомендательной системы: определить "
            "собираемые события, аналитические признаки, правила вывода, способ расчета оценки "
            "рекомендации и вариант дальнейшего развития на основе дерева решений."
        ): (
            "– реализовать экспертную рекомендательную систему: определить события и аналитические "
            "признаки, сформировать правила вывода, рассчитать оценку рекомендации и предусмотреть "
            "дальнейшее развитие на основе дерева решений."
        ),
        (
            "При этом первоначальная версия имела ограничения. Во-первых, задания были представлены "
            "скорее, как локальный список, чем как полноценный семейный процесс. Не было удобного "
            "интерфейса родителя, позволяющего создавать квесты и проверять выполнение. Во-вторых, "
            "интерфейс Unity-клиента оставался функциональным, но недостаточно выразительным: карта "
            "воспринималась как набор кнопок, а не как игровое пространство. В-третьих, адаптивность "
            "под мобильные экраны требовала доработки: некоторые элементы выглядели нормально на ПК, "
            "но плохо масштабировались на телефоне."
        ): (
            "При этом первоначальная версия имела ограничения. Во-первых, задания были представлены "
            "скорее как локальный список, а не как полноценный семейный процесс. Отсутствовал удобный "
            "интерфейс родителя для создания квестов и проверки выполнения. Во-вторых, Unity-интерфейс "
            "оставался функциональным, но недостаточно выразительным: карта воспринималась как набор "
            "кнопок, а не как игровое пространство. В-третьих, требовалась адаптация к мобильным "
            "экранам: элементы, корректно выглядевшие на ПК, плохо масштабировались на телефоне."
        ),
        (
            "К тому же, в первой версии визуальная составляющая выполняла вспомогательную роль. "
            "Для детского продукта этого недостаточно, потому что восприятие игры в значительной "
            "степени строится на ощущении мира: фоне, персонажах, врагах, предметах, иконках, панелях "
            "и состояниях интерфейса. Без стилистического единства даже рабочая механика может "
            "восприниматься как технический прототип, а не как законченный игровой опыт."
        ): (
            "Кроме того, в первой версии визуальная составляющая выполняла вспомогательную роль. "
            "Для детского продукта этого недостаточно, поскольку восприятие игры в значительной "
            "степени определяется целостностью мира: фонами, персонажами, врагами, предметами, "
            "иконками, панелями и состояниями интерфейса. Без стилистического единства даже рабочая "
            "механика воспринимается как технический прототип, а не как законченный игровой опыт."
        ),
    }
    for old, new in replacements.items():
        try:
            replace_exact(document, old, new)
        except ValueError:
            pass

    # Normalize the task list.
    task_paragraphs = []
    in_tasks = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "Для достижения цели были поставлены следующие задачи:":
            in_tasks = True
            continue
        if in_tasks and text.startswith("Практическая значимость"):
            break
        if in_tasks and text.startswith(("–", "-")):
            task_paragraphs.append(paragraph)
    for index, paragraph in enumerate(task_paragraphs):
        text = paragraph.text.strip().replace("–\t", "– ").replace("-\t", "– ")
        if len(text) > 2:
            text = "– " + text[2].lower() + text[3:]
        text = text.rstrip(".;") + ("." if index == len(task_paragraphs) - 1 else ";")
        replace_paragraph_text(paragraph, text)
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.first_line_indent = Cm(-0.63)

    # The analytics and expert system are implemented; geolocation remains prospective.
    heading_changes = {
        "3.12 Цели аналитики и задачи родителя": "3.12 Реализованный аналитический модуль и задачи родителя",
        "3.25 Этапы реализации аналитико-рекомендательного модуля": "3.25 Этапы разработки аналитико-рекомендательного модуля",
        "3.27 Ожидаемые результаты развития проекта": "3.27 Результаты реализации и направления развития",
    }
    for old, new in heading_changes.items():
        try:
            replace_exact(document, old, new)
        except ValueError:
            pass

    section_rewrites = {
        "3.12 Реализованный аналитический модуль и задачи родителя": [
            (
                "В обновленной версии проекта реализован аналитический модуль родительской веб-панели. "
                "Он рассматривает задания не только как отдельные записи Firestore, но и как источник "
                "данных о динамике выполнения: какие категории ребенок чаще завершает, сколько времени "
                "занимает выполнение, как часто требуется повторная проверка и насколько регулярно "
                "забирается награда."
            ),
            (
                "Задача аналитики состоит не в скрытом наблюдении за ребенком, а в поддержке решения "
                "родителя. Интерфейс показывает агрегированные показатели и помогает выбирать посильные "
                "задания, корректировать сложность и избегать чрезмерной нагрузки."
            ),
            (
                "В веб-приложении доступны сводные показатели по дням и категориям, а также рекомендации "
                "для формы создания квеста. Такой подход превращает историю выполнения в практический "
                "инструмент настройки мотивации."
            ),
        ],
        "3.13 Событийная модель и собираемые показатели": [
            (
                "Для аналитики реализована событийная модель. Значимые переходы жизненного цикла "
                "фиксируются отдельными документами questEvents: создание и просмотр задания, отправка "
                "на проверку, одобрение, отклонение, отмена и получение награды."
            ),
            (
                "События создаются в Unity-клиенте и родительском веб-приложении в тех же точках, где "
                "изменяется состояние квеста. Благодаря этому аналитический след связан с реальным "
                "пользовательским действием, а не восстанавливается только по итоговому статусу."
            ),
            (
                "Для каждого события сохраняются идентификаторы семьи, ребенка и квеста, роль "
                "инициатора, клиент-источник, категория, сложность, награда, локальная дата и время, "
                "а при наличии – длительность выполнения и переход статуса. Состав основных событий "
                "представлен в таблице 3.5."
            ),
        ],
        "3.14 Агрегация данных и ежедневные витрины статистики": [
            (
                "На основе questEvents сервис analyticsService формирует документы analyticsDaily и "
                "analyticsByCategory. Первая коллекция содержит показатели по датам, вторая – накопленную "
                "статистику по категориям заданий."
            ),
            (
                "В текущей конфигурации бесплатного тарифа Firebase пересчет выполняется в веб-клиенте "
                "родителя и разрешается правилами Firestore только для ребенка из той же семьи. Исходный "
                "код Cloud Functions подготовлен для переноса агрегации в доверенную серверную среду "
                "после развертывания проекта на тарифе, поддерживающем функции."
            ),
            (
                "Витрины не заменяют исходные события: они содержат подготовленные показатели для "
                "быстрого отображения интерфейса. Основные показатели приведены в таблице 3.6."
            ),
        ],
        "3.17 Постановка задачи рекомендации": [
            (
                "Рекомендательный модуль реализован как средство поддержки родителя при создании нового "
                "квеста. Он не заменяет решение взрослого, а оценивает выбранную категорию и предлагает "
                "сложность и размер награды с учетом накопленной статистики."
            ),
            (
                "Практический результат представлен оценкой от 0 до 100, меткой «рекомендуется», "
                "«нейтрально» или «не рекомендуется», списком причин, рисков и предлагаемых действий. "
                "Родитель может применить совет или сохранить собственные параметры задания."
            ),
            (
                "Модуль использует сведения из questEvents, analyticsDaily и analyticsByCategory. "
                "При отсутствии истории система выдает нейтральную рекомендацию и стандартные параметры, "
                "поэтому форма остается работоспособной и для нового ребенка."
            ),
        ],
        "3.20 Механизм работы экспертной системы в приложении": [
            (
                "Экспертная система запускается при выборе категории в форме создания квеста. "
                "Сервис загружает агрегированную статистику ребенка и формирует набор признаков "
                "RecommendationFeatures."
            ),
            (
                "Механизм вывода последовательно применяет правила базы знаний. Каждое сработавшее "
                "правило изменяет оценку и добавляет объяснение, риск либо практический совет. "
                "Итоговый QuestRecommendation содержит рекомендуемую сложность, награду, числовую "
                "оценку и текстовые основания."
            ),
            (
                "Веб-интерфейс показывает результат рядом с формой создания задания. Дополнительно "
                "реализована коллекция готовых шаблонов, а стандартная награда автоматически зависит "
                "от сложности: ручное изменение опыта и золота вынесено в расширенные настройки."
            ),
            (
                "Unity-клиент получает признак recommended и recommendationScore вместе с квестом и "
                "может визуально выделить рекомендованное задание, не перегружая ребенка аналитическими "
                "подробностями."
            ),
        ],
        "3.24 Интеграция экспертного модуля в Firebase и веб-панель": [
            (
                "Экспертный модуль интегрирован между аналитическими витринами Firestore и формой "
                "создания квеста. Клиентский recommendationEngine применяет типизированный набор правил "
                "и возвращает результат непосредственно интерфейсу."
            ),
            (
                "Результаты пересчета сохраняются в подколлекциях analyticsDaily, analyticsByCategory "
                "и recommendations. Правила Firestore разрешают запись агрегатов только "
                "аутентифицированному родителю той же семьи. Для производственной конфигурации "
                "подготовлены эквивалентные callable-функции, однако их развертывание требует тарифа "
                "Firebase Blaze."
            ),
            (
                "Unity-клиент отправляет события через AnalyticsEventService и продолжает выполнять "
                "основную игровую роль: показывает задания, передает выполнение на проверку и отображает "
                "прогресс. Родительская веб-панель отвечает за анализ и принятие решения."
            ),
        ],
        "3.25 Этапы разработки аналитико-рекомендательного модуля": [
            (
                "Разработка выполнена поэтапно. Сначала добавлена фиксация событий жизненного цикла "
                "квеста в Unity и Web. Затем реализованы агрегаты по дням и категориям и экран аналитики "
                "родительской панели."
            ),
            (
                "На следующем этапе создана база экспертных правил, механизм расчета оценки и формат "
                "объяснимого результата. Рекомендации интегрированы в форму создания квеста вместе с "
                "готовыми шаблонами и автоматическим расчетом награды по сложности."
            ),
            (
                "Дальнейшим этапом остается обучение дерева решений после накопления достаточной и "
                "репрезентативной истории. До этого момента правила экспертной системы обеспечивают "
                "предсказуемое поведение и прозрачное объяснение."
            ),
        ],
        "3.27 Результаты реализации и направления развития": [
            (
                "Реализованный аналитико-рекомендательный контур расширил родительское приложение: "
                "помимо журнала заданий оно показывает динамику выполнения и формирует объяснимые советы "
                "для создания новых квестов."
            ),
            (
                "Экспертная система работает при небольшом объеме данных, поскольку использует "
                "детерминированные правила. По мере накопления истории ее можно дополнить деревом "
                "решений, сохранив экспертные ограничения и текстовые объяснения."
            ),
            (
                "Геолокационные задания остаются направлением дальнейшего развития. Их внедрение "
                "потребует отдельного Unity-сервиса или Android-плагина, явного управления разрешениями "
                "и минимизации сохраняемых координат. Такое разделение позволяет не выдавать проектную "
                "модель геолокации за завершенную функцию."
            ),
        ],
    }

    for heading_text, paragraphs in section_rewrites.items():
        try:
            heading = paragraph_by_text(document, heading_text)
        except ValueError:
            continue
        current = heading._p.getnext()
        for new_text in paragraphs:
            while current is not None and current.tag != qn("w:p"):
                current = current.getnext()
            if current is None:
                break
            p = Paragraph(current, heading._parent)
            if p.style.name.startswith("Heading"):
                break
            replace_paragraph_text(p, new_text)
            current = current.getnext()

    # Replace the conclusion with a factually aligned summary.
    conclusion = paragraph_by_text(document, "ЗАКЛЮЧЕНИЕ")
    source = conclusion._p.getnext()
    new_conclusion = [
        (
            "В ходе выполнения работы исходный Unity-прототип геймифицированного ежедневника "
            "преобразован в распределенную программную систему, объединяющую игровой клиент ребенка, "
            "родительскую веб-панель и облачное хранилище данных."
        ),
        (
            "В Unity доработаны авторизация, адаптивный интерфейс, карта с объектными локациями, "
            "доска приключений, полноэкранный просмотр задания и визуальное начисление опыта и золота. "
            "Графические ресурсы организованы в единую библиотеку и используются при построении "
            "игровых экранов."
        ),
        (
            "Веб-приложение обеспечивает создание семьи и профиля ребенка, выдачу данных для входа, "
            "создание и проверку квестов, использование готовых шаблонов и автоматический расчет "
            "награды по сложности. Синхронизация Unity и Web реализована через Firebase Authentication "
            "и Cloud Firestore с разграничением доступа правилами безопасности."
        ),
        (
            "Дополнительно реализован аналитико-рекомендательный модуль. Он фиксирует события жизненного "
            "цикла заданий, строит агрегаты по дням и категориям и применяет объяснимые экспертные "
            "правила. Родитель получает оценку рекомендации, причины, риски и предлагаемые параметры "
            "нового квеста."
        ),
        (
            "Проверка проекта включала сборку Next.js-приложения, компиляцию функций и Unity-скриптов, "
            "а также прохождение основных пользовательских сценариев. Выявленное ограничение бесплатного "
            "тарифа Firebase учтено: аналитика имеет защищенный клиентский режим пересчета, а серверные "
            "Cloud Functions подготовлены для последующего развертывания."
        ),
        (
            "Поставленная цель достигнута. Разработанная система связывает реальные обязанности ребенка "
            "с игровым прогрессом, предоставляет родителю инструменты управления и анализа и создает "
            "основу для дальнейшего внедрения геолокационных квестов и обучаемой рекомендательной модели."
        ),
    ]
    for text in new_conclusion:
        while source is not None and source.tag != qn("w:p"):
            source = source.getnext()
        if source is None:
            break
        p = Paragraph(source, conclusion._parent)
        if p.style.name == "Heading 1":
            break
        replace_paragraph_text(p, text)
        source = source.getnext()


def add_citations(document: Document) -> None:
    # Fix the six legacy numbers in the geolocation subsection before adding new citations.
    mapping = {"[1]": "[17]", "[2]": "[18]", "[3]": "[19]", "[4]": "[18]", "[5]": "[10]", "[6]": "[12]"}
    geo_markers = (
        "LocationService",
        "Geofencing API",
        "разрешения на доступ к местоположению",
        "Google Play services Geofence",
        "Callable Functions поддерживают",
        "Firestore хорошо подходит",
    )
    for paragraph in document.paragraphs:
        text = paragraph.text
        if any(marker in text for marker in geo_markers):
            for old, new in mapping.items():
                text = text.replace(old, new)
            replace_paragraph_text(paragraph, text)

    citation_targets = [
        ("Одним из способов решения этой проблемы является геймификация.", " [1–3]"),
        ("Первая версия проекта была ориентирована преимущественно", " [4]"),
        ("Unity-клиент представляет собой игровую часть системы", " [5]"),
        ("адаптивность интерфейса", " [6]"),
        ("Input System", " [7]"),
        ("C# используется", " [8]"),
        ("Firebase Authentication используется", " [9]"),
        ("Firestore хранит", " [10]"),
        ("правила Firestore", " [11]"),
        ("Callable Functions", " [12]"),
        ("Веб-приложение родителя отвечает", " [13]"),
        ("компонентный подход React", " [14]"),
        ("типизация TypeScript", " [15]"),
        ("визуальная иерархия", " [16]"),
        ("В Unity базовый доступ к координатам", " [17]"),
        ("Android Geofencing API", " [18]"),
        ("разрешения на доступ к местоположению", " [19]"),
        ("дерева решений", " [20]"),
    ]
    intro_seen = False
    for needle, suffix in citation_targets:
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text == "ВВЕДЕНИЕ":
                intro_seen = True
                continue
            if not intro_seen or text == "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ":
                continue
            if (
                not paragraph.style.name.startswith("Heading")
                and needle.lower() in text.lower()
                and suffix.strip() not in text
            ):
                paragraph.add_run(suffix)
                break
        intro_seen = False

    def first_body_paragraph_after(heading_text: str) -> Paragraph:
        heading = paragraph_by_text(document, heading_text)
        current = heading._p.getnext()
        while current is not None:
            if current.tag == qn("w:p"):
                paragraph = Paragraph(current, heading._parent)
                if paragraph.text.strip():
                    return paragraph
            current = current.getnext()
        raise ValueError(f"No body paragraph after {heading_text}")

    guaranteed = {
        "1.6 Unity-клиент ребенка": ("[5]", "[8]"),
        "1.7 Родительское веб-приложение": ("[13]", "[14]", "[15]"),
        "1.10 Визуальная система и адаптивный интерфейс": ("[16]",),
    }
    for heading_text, citations in guaranteed.items():
        paragraph = first_body_paragraph_after(heading_text)
        for citation in citations:
            if citation not in paragraph.text:
                paragraph.add_run(" " + citation)

    # Remove duplicate citations created from legacy numbers.
    for paragraph in document.paragraphs:
        text = paragraph.text
        for citation in re.findall(r"\[\d+\]", text):
            if text.count(citation) > 1:
                first = text.find(citation)
                prefix = text[: first + len(citation)]
                suffix = text[first + len(citation):].replace(" " + citation, "").replace(citation, "")
                text = prefix + suffix
        if text != paragraph.text:
            replace_paragraph_text(paragraph, text)


def remove_malformed_table(document: Document) -> None:
    for table in list(document.tables):
        if len(table.rows) == 70 and len(table.columns) == 6:
            delete_element(table._tbl)


def caption_for_table(table: Table) -> Paragraph | None:
    prev = table._tbl.getprevious()
    while prev is not None:
        if prev.tag == qn("w:p"):
            p = Paragraph(prev, table._parent)
            if p.text.strip():
                return p
        prev = prev.getprevious()
    return None


def add_and_fix_captions(document: Document) -> None:
    table_caption_data = [
        ("Таблица 3.1 – Поля расширенной модели задания",
         "Состав полей расширенной модели задания представлен в таблице 3.1."),
        ("Таблица 3.2 – Сравнение способов геолокационной проверки",
         "Сравнение способов геолокационной проверки представлено в таблице 3.2."),
        ("Таблица 3.3 – Поля геолокационного задания",
         "Состав полей геолокационного задания представлен в таблице 3.3."),
        ("Таблица 3.4 – Сценарии проверки геолокационных квестов",
         "Основные сценарии проверки геолокационных квестов представлены в таблице 3.4."),
    ]
    tables = list(document.tables)
    for index, (caption, reference_sentence) in enumerate(table_caption_data):
        table = tables[index]
        caption_p = insert_paragraph_before_element(document, table._tbl, caption)
        set_compact_format(caption_p, WD_ALIGN_PARAGRAPH.LEFT, 14, None, 0, 1.0)
        previous = caption_p._p.getprevious()
        while previous is not None and previous.tag != qn("w:p"):
            previous = previous.getprevious()
        if previous is not None:
            prev_p = Paragraph(previous, caption_p._parent)
            if reference_sentence not in prev_p.text:
                prev_p.add_run(" " + reference_sentence)

    analytics_events_table = list(document.tables)[4]
    analytics_caption = caption_for_table(analytics_events_table)
    if analytics_caption is None or not analytics_caption.text.strip().startswith("Таблица"):
        caption_p = insert_paragraph_before_element(
            document,
            analytics_events_table._tbl,
            "Таблица 3.5 – События жизненного цикла квеста",
        )
        set_compact_format(caption_p, WD_ALIGN_PARAGRAPH.LEFT, 14, None, 0, 1.0)
        caption_p.paragraph_format.keep_with_next = True

    caption_replacements = {
        "Таблица 7.1 - События жизненного цикла квеста": "Таблица 3.5 – События жизненного цикла квеста",
        "Таблица 7.2 - Основные аналитические показатели": "Таблица 3.6 – Основные аналитические показатели",
        "Таблица 8.1 - Признаки рекомендательной системы": "Таблица 3.7 – Признаки рекомендательной системы",
        "Таблица 8.2 - Примеры правил экспертной системы": "Таблица 3.8 – Примеры правил экспертной системы",
    }
    figure_replacements = {
        "Рисунок 1.1 - Общая архитектура системы": "Рисунок 1.1 – Общая архитектура системы",
        "Рисунок 1.2 - Структура Firestore": "Рисунок 1.2 – Структура Firestore",
        "Рисунок 1.3 - Жизненный цикл задания": "Рисунок 1.3 – Жизненный цикл задания",
        "Рисунок 3.1 - Схема аналитико-рекомендательного контура":
            "Рисунок 3.1 – Схема аналитико-рекомендательного контура",
        "Рисунок 3.2 - Структура экспертной системы рекомендаций":
            "Рисунок 3.2 – Структура экспертной системы рекомендаций",
    }
    for old, new in {**caption_replacements, **figure_replacements}.items():
        try:
            replace_exact(document, old, new)
        except ValueError:
            pass

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Таблица "):
            set_compact_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT, 14, None, 0, 1.0)
            paragraph.paragraph_format.keep_with_next = True
        elif text.startswith("Рисунок "):
            set_compact_format(paragraph, WD_ALIGN_PARAGRAPH.CENTER, 14, None, 0, 1.0)


def format_tables(document: Document) -> None:
    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    set_compact_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT, 12, row_index == 0, 0, 1.0)
                    paragraph.paragraph_format.keep_together = True
        if table.rows:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)


REFERENCES = [
    "1. Deterding, S. From game design elements to gamefulness: defining gamification / "
    "S. Deterding, D. Dixon, R. Khaled, L. Nacke // Proceedings of the 15th International "
    "Academic MindTrek Conference. – New York : ACM, 2011. – P. 9–15. – "
    "DOI 10.1145/2181037.2181040.",
    "2. Kapp, K. M. The Gamification of Learning and Instruction: Game-based Methods and "
    "Strategies for Training and Education / K. M. Kapp. – San Francisco : Pfeiffer, 2012. – "
    "336 p. – ISBN 978-1-118-09634-5.",
    "3. Werbach, K. For the Win: How Game Thinking Can Revolutionize Your Business / "
    "K. Werbach, D. Hunter. – Philadelphia : Wharton Digital Press, 2012. – 148 p. – "
    "ISBN 978-1-61363-023-5.",
    "4. Шелл, Дж. Геймдизайн. Как создать игру, в которую будут играть все / Дж. Шелл. – "
    "Москва : Альпина Паблишер, 2019. – 640 с. – ISBN 978-5-9614-2512-3.",
    "5. Unity 6.4 User Manual // Unity Documentation : [сайт]. – URL: "
    "https://docs.unity3d.com/6000.4/Documentation/Manual/ "
    "(дата обращения: 12.06.2026).",
    "6. Designing UI for Multiple Resolutions // Unity Documentation : [сайт]. – URL: "
    "https://docs.unity3d.com/Manual/HOWTO-UIMultiResolution.html "
    "(дата обращения: 12.06.2026).",
    "7. Input System // Unity Documentation : [сайт]. – URL: "
    "https://docs.unity3d.com/Packages/com.unity.inputsystem@1.14/manual/ "
    "(дата обращения: 12.06.2026).",
    "8. Асинхронное программирование с использованием async и await // Microsoft Learn : "
    "[сайт]. – URL: https://learn.microsoft.com/dotnet/csharp/asynchronous-programming/ "
    "(дата обращения: 12.06.2026).",
    "9. Firebase Authentication // Firebase : [сайт]. – URL: "
    "https://firebase.google.com/docs/auth (дата обращения: 12.06.2026).",
    "10. Cloud Firestore Data Model // Firebase : [сайт]. – URL: "
    "https://firebase.google.com/docs/firestore/data-model (дата обращения: 12.06.2026).",
    "11. Cloud Firestore Security Rules // Firebase : [сайт]. – "
    "URL: https://firebase.google.com/docs/firestore/security/ "
    "(дата обращения: 12.06.2026).",
    "12. Callable Functions // Firebase : [сайт]. – URL: "
    "https://firebase.google.com/docs/functions/callable (дата обращения: 12.06.2026).",
    "13. App Router // Next.js : [сайт]. – URL: "
    "https://nextjs.org/docs/app (дата обращения: 12.06.2026).",
    "14. React // [сайт]. – URL: https://react.dev/learn "
    "(дата обращения: 12.06.2026).",
    "15. TypeScript Handbook // TypeScript : [сайт]. – URL: "
    "https://www.typescriptlang.org/docs/handbook/ (дата обращения: 12.06.2026).",
    "16. Norman, D. A. The Design of Everyday Things / D. A. Norman. – Revised and expanded ed. – "
    "New York : Basic Books, 2013. – 368 p. – ISBN 978-0-465-05065-9.",
    "17. LocationService.Start // Unity : [сайт]. – URL: "
    "https://docs.unity3d.com/6000.4/Documentation/ScriptReference/LocationService.Start.html "
    "(дата обращения: 12.06.2026).",
    "18. Geofencing // Android Developers : [сайт]. – URL: "
    "https://developer.android.com/develop/sensors-and-location/location/geofencing "
    "(дата обращения: 12.06.2026).",
    "19. Location Permissions // Android Developers : [сайт]. – URL: "
    "https://developer.android.com/develop/sensors-and-location/location/permissions "
    "(дата обращения: 12.06.2026).",
    "20. Decision Trees // scikit-learn : [сайт]. – URL: "
    "https://scikit-learn.org/stable/modules/tree.html (дата обращения: 12.06.2026).",
]


def replace_references(document: Document) -> None:
    heading = paragraph_by_text(document, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    current = heading._p.getnext()
    while current is not None and current.tag != qn("w:sectPr"):
        nxt = current.getnext()
        delete_element(current)
        current = nxt
    for reference in REFERENCES:
        p = document.add_paragraph(reference)
        set_normal_format(p, True)


def global_typography(document: Document) -> None:
    intro_reached = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "ВВЕДЕНИЕ":
            intro_reached = True
        if not intro_reached:
            continue
        if not text:
            set_compact_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT, 14, None, 0, 1.5)
            continue
        if paragraph.style.name.startswith("Heading"):
            continue
        if text.startswith(("Рисунок ", "Таблица ")):
            continue
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            continue
        set_normal_format(paragraph, True)

def add_heading_blank_lines(document: Document) -> None:
    for paragraph in list(document.paragraphs):
        if not paragraph.style.name.startswith("Heading"):
            continue
        nxt = paragraph._p.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            next_p = Paragraph(nxt, paragraph._parent)
            if not next_p.text.strip():
                continue
        blank = insert_paragraph_after(paragraph)
        set_compact_format(blank, WD_ALIGN_PARAGRAPH.LEFT, 14, None, 0, 1.5)


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_structural_heading_display(document: Document) -> None:
    replacements = {
        "РЕФЕРАТ": "Реферат",
        "СОДЕРЖАНИЕ": "Содержание",
        "ВВЕДЕНИЕ": "Введение",
        "ЗАКЛЮЧЕНИЕ": "Заключение",
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ": "Список использованных источников",
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text not in replacements:
            continue
        replace_paragraph_text(paragraph, replacements[text])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.keep_with_next = True
        set_paragraph_font(paragraph, 14, True)
        for run in paragraph.runs:
            run.font.all_caps = True


def fix_quotes_and_dashes(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading") or paragraph.text.strip().startswith(("http",)):
            continue
        text = paragraph.text
        text = re.sub(r"\s-\s", " – ", text)
        text = text.replace("В итоге,", "В итоге")
        text = text.replace('"Рекомендация системы"', "«Рекомендация системы»")
        text = text.replace('"Совет дня"', "«Совет дня»")
        text = text.replace('"рекомендуется"', "«рекомендуется»")
        if text != paragraph.text:
            replace_paragraph_text(paragraph, text)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document(SOURCE)
    configure_styles(document)
    configure_page(document)
    intro = remove_front_matter(document)
    build_front_matter(document, intro)
    normalize_headings(document)
    edit_content(document)
    add_citations(document)
    remove_malformed_table(document)
    add_and_fix_captions(document)
    replace_references(document)
    fix_quotes_and_dashes(document)
    format_tables(document)
    global_typography(document)
    normalize_headings(document)
    add_heading_blank_lines(document)
    set_structural_heading_display(document)
    set_update_fields(document)
    document.core_properties.title = (
        "Разработка мобильного приложения-ежедневника с использованием метода геймификации"
    )
    document.core_properties.author = "И. Д. Ерюшев"
    document.core_properties.subject = "Выпускная квалификационная работа"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
